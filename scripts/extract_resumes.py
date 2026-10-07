"""Extract text from resumes/raw/* -> resumes/text/*.txt, dedupe identical text, write manifest."""
import hashlib, json, re, sys
from pathlib import Path
import docx
from pypdf import PdfReader

RAW = Path(__file__).resolve().parent.parent / "resumes" / "raw"
OUT = RAW.parent / "text"
OUT.mkdir(exist_ok=True)

def text_of(p: Path) -> str:
    if p.suffix == ".docx":
        d = docx.Document(p)
        lines = [x.text for x in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                cells = []
                for c in row.cells:
                    if c.text not in cells:
                        cells.append(c.text)
                lines.append(" | ".join(cells))
        return "\n".join(l for l in lines if l.strip())
    if p.suffix == ".pdf":
        if p.stat().st_size == 0:
            raise ValueError("0-byte file in Drive")
        t = "\n".join((pg.extract_text() or "") for pg in PdfReader(p).pages)
        lines = [l for l in t.splitlines() if l.strip()]
        if lines and sum(map(len, lines)) / len(lines) < 15:  # one word per line -> re-flow
            t = re.sub(r"[ \t]*\n[ \t]*", " ", t)
            t = re.sub(r"\s+([•·|])\s+", r"\n\1 ", t)
        return t
    return p.read_text(errors="replace")

def norm(t):
    return re.sub(r"\s+", " ", t).strip().lower()

manifest, seen = [], {}
for p in sorted(RAW.iterdir()):
    entry = {"file": p.name, "bytes": p.stat().st_size}
    try:
        t = text_of(p)
        if len(t.strip()) < 50:
            raise ValueError("no extractable text")
        if t.lstrip().lower().startswith("<!doctype html") and "resume" not in t.lower()[:3000]:
            raise ValueError("not a resume (HTML tool-valuation page)")
        h = hashlib.sha1(norm(t).encode()).hexdigest()[:12]
        entry.update(text_sha=h, chars=len(t))
        if h in seen:
            entry["duplicate_of"] = seen[h]
        else:
            seen[h] = p.name
            (OUT / (p.stem + ".txt")).write_text(t)
    except Exception as e:
        entry["error"] = str(e)
    manifest.append(entry)
json.dump(manifest, open(RAW.parent / "manifest.json", "w"), indent=1)
for m in manifest:
    print(f"{m['file'][:75]:75s} {m.get('chars','-'):>6} {m.get('text_sha','')} {('DUP of '+m['duplicate_of'][:40]) if 'duplicate_of' in m else ''} {m.get('error','')}")
print("unique texts:", len(seen))
